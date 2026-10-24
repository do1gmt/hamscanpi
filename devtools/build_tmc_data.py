#!/usr/bin/env python3
"""
TMC-Daten für HamScanPi aus den Originaldateien der BASt erzeugen.

Quelle: Bundesanstalt für Straßenwesen (BASt), „Location Code List“ LCL 22.0 (Deutschland, LTN 1)
und „Event Code List“ ECL 4.01 (deutsch), Lizenz CC BY 4.0
(https://creativecommons.org/licenses/by/4.0/). Die Daten wurden in ein kompaktes JSON-Format
umgewandelt und auf die für HamScanPi nötigen Felder gekürzt (Punkte, Straßen, Abschnitte, Gebiete,
Ereignis- und Zusatztexte); inhaltlich nicht verändert.

  python3 devtools/build_tmc_data.py <entpackter Ordner LCL22_deutsch>
  → share/tmc/lcl_de.json.gz, share/tmc/events_de.json.gz
"""
import csv
import gzip
import json
import sys
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "share" / "tmc"


def dat(folder, name):
    p = next(folder.rglob(f"{name}.DAT"))
    with open(p, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def num(x):
    try:
        return int(x)
    except (TypeError, ValueError):
        return None


def coord(x):
    try:
        return round(int(x) / 1e5, 5)
    except (TypeError, ValueError):
        return None


def main(src):
    src = Path(src)
    names = {r["NID"]: r["NAME"] for r in dat(src, "NAMES")}
    sub = {(r["CLASS"], r["TCD"], r["STCD"]): r["SNATDESC"] for r in dat(src, "SUBTYPES")}
    poff = {r["LCD"]: (num(r["NEG_OFF_LCD"]), num(r["POS_OFF_LCD"])) for r in dat(src, "POFFSETS")}
    soff = {r["LCD"]: (num(r["NEG_OFF_LCD"]), num(r["POS_OFF_LCD"])) for r in dat(src, "SOFFSETS")}
    ltn = num(dat(src, "LOCATIONDATASETS")[0]["TABCD"])
    roads = {}
    for r in dat(src, "ROADS"):
        roads[r["LCD"]] = [r["ROADNUMBER"], names.get(r["N1ID"], ""), names.get(r["N2ID"], ""), names.get(r["RNID"], "")]
    segs = {}
    for r in dat(src, "SEGMENTS"):
        n, p = soff.get(r["LCD"], (None, None))
        segs[r["LCD"]] = [r["ROADNUMBER"], names.get(r["N1ID"], ""), names.get(r["N2ID"], ""), num(r["ROA_LCD"]), n, p]
    pts = {}
    for r in dat(src, "POINTS"):
        n, p = poff.get(r["LCD"], (None, None))
        road = num(r["ROA_LCD"])
        if road is None and r["SEG_LCD"] in segs:
            road = segs[r["SEG_LCD"]][3]
        pts[r["LCD"]] = [road, num(r["SEG_LCD"]), names.get(r["N1ID"], ""), names.get(r["N2ID"], ""),
                         r["JUNCTIONNUMBER"] or "", sub.get(("P", r["TCD"], r["STCD"]), ""), n, p,
                         coord(r["YCOORD"]), coord(r["XCOORD"])]
    areas = {}
    for f in ("ADMINISTRATIVEAREA", "OTHERAREAS"):
        for r in dat(src, f):
            areas[r["LCD"]] = names.get(r["NID"], "")
    data = {"source": "BASt Location Code List LCL 22.0 (CC BY 4.0)", "ltn": ltn,
            "fields": {"points": "road, segment, name1, name2, junction, type, neg_off, pos_off, lat, lon",
                       "roads": "number, name1, name2, roadname", "segments": "number, name1, name2, road, neg_off, pos_off"},
            "points": pts, "roads": roads, "segments": segs, "areas": areas}
    OUT.mkdir(parents=True, exist_ok=True)
    with gzip.open(OUT / "lcl_de.json.gz", "wt", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))

    import openpyxl
    wb = openpyxl.load_workbook(next(src.rglob("ECL*.xlsx")), read_only=True)
    ev, sup = {}, {}
    for sheet, col_txt, col_q1, col_qn, col_code, col_en in (("ForecastEvent List", 2, 3, 4, 5, 1), ("Event List", 3, 4, 5, 6, 1)):
        for row in wb[sheet].iter_rows(values_only=True):
            code = row[col_code] if len(row) > col_code else None
            if isinstance(code, (int, float)) and row[col_txt]:
                ev[int(code)] = {"de": str(row[col_txt]).strip(), "q1": (row[col_q1] or "").strip() if isinstance(row[col_q1], str) else "",
                                 "qn": (row[col_qn] or "").strip() if isinstance(row[col_qn], str) else "",
                                 "en": (row[col_en] or "").strip() if isinstance(row[col_en], str) else ""}
    for row in wb["Supplementary Information List"].iter_rows(values_only=True):
        if len(row) > 3 and isinstance(row[3], (int, float)) and row[2]:
            sup[int(row[3])] = str(row[2]).strip()
    with gzip.open(OUT / "events_de.json.gz", "wt", encoding="utf-8") as f:
        json.dump({"source": "BASt Event Code List ECL 4.01 (CC BY 4.0)", "events": ev, "supplementary": sup},
                  f, ensure_ascii=False, separators=(",", ":"))
    print(f"Punkte {len(pts)}, Straßen {len(roads)}, Abschnitte {len(segs)}, Gebiete {len(areas)}, "
          f"Ereignisse {len(ev)}, Zusatztexte {len(sup)}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else ".")
